"""Inspect the Phase 1 workbook — fills, dimensions, freeze panes, sample cells."""
from openpyxl import load_workbook

PATH = r"C:\Users\iAmiK\Downloads\Dan_Phase1_Integrated_Workbook.xlsx"
EXPECTED_BG_COLORS = {
    '00000000', 'FF1A0B2E', 'FF2D1B4E', 'FF3D2A5F',
    'FF4A2C7A', 'FF5E3D94', 'FF7B4FB8',
    'FF9D6FE8', 'FFC8A8F0', 'FFE8C547',
    'FF6BCB77', 'FFFFD93D', 'FFF5F0FF',
}

def main():
    wb = load_workbook(PATH)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        used_rows = ws.max_row
        used_cols = ws.max_column
        # sample: count unfilled cells in used range (max 200 rows checked to be fast)
        unfilled = 0
        checked = 0
        sample_colors = set()
        for r in range(1, min(used_rows, 200) + 1):
            for c in range(1, used_cols + 1):
                cell = ws.cell(r, c)
                fill = cell.fill
                checked += 1
                if fill is None or fill.fgColor is None:
                    unfilled += 1
                    continue
                rgb = fill.fgColor.rgb
                if rgb in (None, '00000000'):
                    unfilled += 1
                else:
                    sample_colors.add(rgb)
        freeze = ws.freeze_panes
        merges = len(ws.merged_cells.ranges)
        print(f'\n[{sheet_name}]')
        print(f'  Dimensions: {used_rows} rows × {used_cols} cols')
        print(f'  Cells checked (first 200 rows): {checked}')
        print(f'  Cells without fill: {unfilled}')
        print(f'  Unique fill colors: {len(sample_colors)}')
        print(f'  Freeze panes: {freeze or "(none)"}')
        print(f'  Merged ranges: {merges}')
        # Column widths
        widths = [ws.column_dimensions[chr(64+i)].width for i in range(1, used_cols + 1)]
        print(f'  Column widths: {widths}')

if __name__ == '__main__':
    main()
