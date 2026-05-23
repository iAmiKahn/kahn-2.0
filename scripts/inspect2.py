"""Check for truly unfilled cells (outside merged ranges)."""
from openpyxl import load_workbook

PATH = r"C:\Users\iAmiK\Downloads\Dan_Phase1_Integrated_Workbook.xlsx"

def main():
    wb = load_workbook(PATH)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        used_rows = ws.max_row
        used_cols = ws.max_column
        # Build a set of coords that are inside a merged range (excluding top-left)
        merged_inner = set()
        for mr in ws.merged_cells.ranges:
            for row in range(mr.min_row, mr.max_row + 1):
                for col in range(mr.min_col, mr.max_col + 1):
                    if (row, col) != (mr.min_row, mr.min_col):
                        merged_inner.add((row, col))

        truly_unfilled = []
        for r in range(1, used_rows + 1):
            for c in range(1, used_cols + 1):
                if (r, c) in merged_inner:
                    continue
                cell = ws.cell(r, c)
                fill = cell.fill
                if fill is None or fill.fgColor is None:
                    truly_unfilled.append((r, c))
                    continue
                rgb = fill.fgColor.rgb
                if rgb in (None, '00000000'):
                    truly_unfilled.append((r, c))

        print(f'[{sheet_name}] {used_rows}×{used_cols}  truly unfilled (NOT in merged inner): {len(truly_unfilled)}')
        if truly_unfilled[:5]:
            print(f'  samples: {truly_unfilled[:5]}')

if __name__ == '__main__':
    main()
