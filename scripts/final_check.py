"""Final sanity check: daily grid content, no PER-DAY column, gold callout, tab guide."""
from openpyxl import load_workbook

PATH = r"C:\Users\iAmiK\Downloads\Dan_Phase1_Integrated_Workbook.xlsx"

def main():
    wb = load_workbook(PATH)

    # 1. INGREDIENTS: confirm daily grid exists by looking for "DAILY INGREDIENT GRID" text
    print('--- INGREDIENTS & SHOPPING ---')
    ws = wb['INGREDIENTS & SHOPPING']
    grid_row = None
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and 'DAILY INGREDIENT GRID' in v:
                grid_row = r
                print(f'  "DAILY INGREDIENT GRID" banner found at row {r}, col {c}')
                break
        if grid_row: break

    # Confirm no PER-DAY column in top section (should only find "PER-DAY" if present)
    per_day_found = False
    for r in range(1, 90):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().upper() == 'PER-DAY':
                per_day_found = True
                print(f'  !! PER-DAY column header found at {ws.cell(r,c).coordinate}')
    if not per_day_found:
        print('  PER-DAY column confirmed removed.')

    # Spot-check daily grid: look for chicken breast row
    print('\n  Daily grid spot checks:')
    if grid_row is not None:
        for target in ['Chicken breast', 'Ground beef 93/7', 'Greek yogurt']:
            for r in range(grid_row, ws.max_row + 1):
                name = ws.cell(r, 2).value
                if isinstance(name, str) and target in name:
                    # Read MON..SUN..WEEKLY from cols 3-10
                    mon = ws.cell(r, 3).value
                    tue = ws.cell(r, 4).value
                    wed = ws.cell(r, 5).value
                    thu = ws.cell(r, 6).value
                    fri = ws.cell(r, 7).value
                    sat = ws.cell(r, 8).value
                    sun = ws.cell(r, 9).value
                    wkly = ws.cell(r, 10).value
                    print(f'    {name}: MON={mon} TUE={tue} WED={wed} THU={thu} FRI={fri} SAT={sat} SUN={sun} | WEEKLY={wkly}')
                    break

    # 2. START HERE: confirm gold callout + tab guide
    print('\n--- START HERE ---')
    ws = wb['START HERE']
    gold_found = False
    tab_guide_found = False
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                if 'scale is the law' in v.lower():
                    gold_found = True
                    print(f'  Gold callout found at {ws.cell(r,c).coordinate}')
                if 'TAB GUIDE' in v.upper():
                    tab_guide_found = True
                    print(f'  Tab guide found at {ws.cell(r,c).coordinate}')
    if not gold_found:
        print('  !! gold callout MISSING')
    if not tab_guide_found:
        print('  !! tab guide MISSING')

    # 3. Freeze panes
    print('\n--- Freeze Panes ---')
    mp = wb['WEEKLY MEAL PLAN'].freeze_panes
    fs = wb['FOOD SWAPS'].freeze_panes
    print(f'  WEEKLY MEAL PLAN: {mp}  ({"ok" if mp == "A6" else "MISSING"})')
    print(f'  FOOD SWAPS: {fs}  ({"ok" if fs == "A8" else "MISSING"})')

    # 4. Language scan
    print('\n--- Language Scan ---')
    flags = []
    bad = ['Dan\'s job', 'Brett\'s job', 'your responsibility', 'you\'ll need']
    for sheet_name in wb.sheetnames:
        ws_ = wb[sheet_name]
        for row in ws_.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str):
                    for b in bad:
                        if b in v:
                            flags.append(f'{sheet_name}!{cell.coordinate}: "{b}" in "{v[:60]}"')
    if flags:
        for f in flags: print(f'  !! {f}')
    else:
        print('  clean.')

    # 5. Totals per sheet (dimensions)
    print('\n--- Sheet Dimensions ---')
    for s in wb.sheetnames:
        ws_ = wb[s]
        print(f'  {s}: {ws_.max_row} × {ws_.max_column}  (merges: {len(ws_.merged_cells.ranges)})')

if __name__ == '__main__':
    main()
