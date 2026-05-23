"""Run Dan's literal unfilled-cell check from the correction brief."""
from openpyxl import load_workbook

PATH = r"C:\Users\iAmiK\Downloads\Dan_Phase1_Integrated_Workbook.xlsx"

def main():
    wb = load_workbook(PATH)
    total = 0
    for tab in wb.sheetnames:
        ws = wb[tab]
        unfilled = 0
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                c = ws.cell(row=row, column=col)
                fill = c.fill.start_color.value if c.fill.start_color else None
                if fill in (None, "00000000", "FFFFFFFF", "00FFFFFF"):
                    unfilled += 1
        print(f"{tab}: {unfilled} unfilled cells")
        total += unfilled
    print(f'\nTOTAL: {total} unfilled cells across all tabs')

if __name__ == '__main__':
    main()
